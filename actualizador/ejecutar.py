
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import argparse
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo
from itertools import zip_longest
import pandas as pd
from openpyxl import load_workbook

from actualizador.databricks import obtener_ordenes, obtener_facturacion
from actualizador.fuentes import preparar_fuentes
from m365.config import (
    SITE_NAME,
    DRIVE_NAME,
    INPUT_FOLDER,
    PROJECT_FOLDER,
    TALLER_FILE,
    CLIENTES_FILE,
)
from m365.sharepoint import SharePoint
from m365.excel import read_excel


SERVICIOS_FILE = "nivel de servicio.xlsx"
MODELOS_FILE = "modelos vehiculos.xlsx"
BACKUP_FOLDER = f"{PROJECT_FOLDER}/backup"

COLUMNAS_CLIENTES = [
    "nro documento", "nombre", "telefono", "correo"
]

COLUMNAS_TEXTO = {
    TALLER_FILE: {"Id doc", "doc ref", "placa", "vin",
                  "modelo", "año", "sede", "asesor",
                  "nro documento", "servicio1", "servicio2"},
    CLIENTES_FILE: set(COLUMNAS_CLIENTES),
    MODELOS_FILE: {"Material", "Modelo", "Año"},
}


def texto(valor) -> str:
    if pd.isna(valor):
        return ""
    return str(valor).strip()


def conectar_sharepoint():
    sp = SharePoint()
    site = sp.get_site(SITE_NAME)

    if site is None:
        raise RuntimeError("No se encontró el sitio SharePoint.")

    drive = sp.get_drive(site["id"], DRIVE_NAME)

    if drive is None:
        raise RuntimeError("No se encontró la biblioteca.")

    return sp, drive["id"]


def descargar(sp, drive_id, nombre):
    ruta = f"{INPUT_FOLDER}/{nombre}"
    archivo = sp.download_file_by_path(drive_id, ruta)

    contenido = archivo.getvalue()
    datos = read_excel(
        BytesIO(contenido),
        dtype="string",
    )

    return contenido, datos


def filtrar_clientes_nuevos(
    historico: pd.DataFrame,
    propuestas: pd.DataFrame,
) -> pd.DataFrame:
    """
    No elimina registros históricos.

    Descarta propuestas idénticas al histórico y evita
    repetir filas idénticas dentro del lote nuevo.
    """
    if propuestas.empty:
        return propuestas.copy()

    def normalizar(df):
        resultado = df[COLUMNAS_CLIENTES].copy()

        for columna in COLUMNAS_CLIENTES:
            resultado[columna] = (
                resultado[columna]
                .astype("string")
                .fillna("")
                .str.strip()
            )

        return resultado

    antiguos = normalizar(historico)
    nuevos = normalizar(propuestas)

    existentes = set(
        antiguos.itertuples(index=False, name=None)
    )

    conservar = []
    vistos = set()

    for fila in nuevos.itertuples(index=False, name=None):
        if not fila[0]:
            continue

        if fila in existentes or fila in vistos:
            continue

        conservar.append(fila)
        vistos.add(fila)

    return pd.DataFrame(
        conservar,
        columns=COLUMNAS_CLIENTES,
    )


def agregar_filas_excel(
    contenido_original: bytes,
    nuevas: pd.DataFrame,
    nombre_archivo: str,
) -> BytesIO:
    """
    Agrega filas al libro original conservando las hojas
    existentes y su formato general.
    """
    wb = load_workbook(BytesIO(contenido_original))
    ws = wb.worksheets[0]

    encabezados = {
        str(celda.value).strip(): celda.column
        for celda in ws[1]
        if celda.value is not None
    }

    faltantes = set(nuevas.columns) - set(encabezados)

    if faltantes:
        raise ValueError(
            f"{nombre_archivo}: columnas ausentes: "
            f"{sorted(faltantes)}"
        )

    # Localizar la última fila con datos reales.
    ultima_fila_datos = 1

    for fila in ws.iter_rows(min_row=2):
        if any(
            celda.value is not None
            and str(celda.value).strip() != ""
            for celda in fila
        ):
            ultima_fila_datos = fila[0].row

    for valores in nuevas.to_dict("records"):
        ultima_fila_datos += 1
        fila_destino = ultima_fila_datos

        for columna, valor in valores.items():
            celda = ws.cell(
                row=fila_destino,
                column=encabezados[columna],
            )

            if pd.isna(valor):
                continue

            if isinstance(valor, pd.Timestamp):
                celda.value = valor.to_pydatetime()
                celda.number_format = "dd/mm/yyyy"
            elif columna in COLUMNAS_TEXTO.get(
                nombre_archivo, set()
            ):
                celda.value = str(valor)
                celda.number_format = "@"
            else:
                celda.value = valor

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)

    return salida


def validar_incremento(
    original: pd.DataFrame,
    resultado: pd.DataFrame,
    nuevas: pd.DataFrame,
    nombre: str,
) -> None:
    if len(resultado) != len(original) + len(nuevas):
        raise RuntimeError(
            f"Inconsistencia de filas en {nombre}."
        )

    if list(original.columns) != list(resultado.columns):
        raise RuntimeError(
            f"Se modificó la estructura de {nombre}."
        )

    pd.testing.assert_frame_equal(
        resultado.iloc[:len(original)].reset_index(drop=True),
        original.reset_index(drop=True),
        check_dtype=False,
    )


def cargar_excel_desde_bytes(contenido: bytes) -> pd.DataFrame:
    return read_excel(BytesIO(contenido), dtype="string")


def verificar_archivo(
    contenido: bytes,
    filas_esperadas: int,
    nombre: str,
) -> None:
    df = cargar_excel_desde_bytes(contenido)

    if len(df) != filas_esperadas:
        raise RuntimeError(
            f"{nombre}: se esperaban {filas_esperadas} "
            f"filas, pero se encontraron {len(df)}."
        )


def respaldar(sp, drive_id, originales, archivos):
    """
    Crea o sobrescribe un backup rotativo por cada archivo
    que será modificado.

    Verifica que el contenido respaldado sea idéntico
    al original antes de autorizar la publicación.
    """
    for nombre in archivos:
        nombre_backup = (
            f"{nombre.removesuffix('.xlsx')}_backup.xlsx"
        )

        ruta = f"{BACKUP_FOLDER}/{nombre_backup}"

        print(f"💾 Respaldando {nombre}...")

        sp.upload_file_by_path(
            drive_id=drive_id,
            file_path=ruta,
            file_content=BytesIO(originales[nombre]),
        )

        copia = sp.download_file_by_path(
            drive_id=drive_id,
            file_path=ruta,
        ).getvalue()

        if copia != originales[nombre]:
            raise RuntimeError(
                f"Backup de {nombre} no coincide "
                "con el archivo original."
            )

        print(f"   ✅ Backup verificado: {nombre_backup}")


def _celdas_con_contenido(hoja):
    """
    Recorre las celdas con valores o fórmulas.

    Ignora las celdas vacías y las filas que únicamente
    contienen formato.
    """
    for fila in hoja.iter_rows():
        for celda in fila:
            valor = celda.value

            if valor is None or valor == "":
                continue

            yield (
                celda.coordinate,
                valor,
            )


def verificar_publicacion_excel(
    original: bytes,
    descargado: bytes,
    nombre_archivo: str,
) -> None:
    """
    Comprueba que un Excel publicado conserva el
    contenido del Excel preparado.

    Primero compara bytes. Si difieren, compara los
    nombres de las hojas, las celdas y sus valores.

    Las fórmulas se comparan como fórmulas, no por
    sus resultados calculados.
    """
    if original == descargado:
        print(
            f"   ✅ Verificación binaria: {nombre_archivo}"
        )
        return

    print(
        f"   ℹ️ Diferencia binaria en {nombre_archivo}. "
        "Verificando contenido del Excel..."
    )

    libro_original = load_workbook(
        BytesIO(original),
        read_only=True,
        data_only=False,
    )

    libro_descargado = None

    try:
        libro_descargado = load_workbook(
            BytesIO(descargado),
            read_only=True,
            data_only=False,
        )

        if libro_original.sheetnames != libro_descargado.sheetnames:
            raise RuntimeError(
                f"{nombre_archivo}: las hojas no coinciden. "
                f"Original: {libro_original.sheetnames}. "
                f"Publicado: {libro_descargado.sheetnames}."
            )

        for nombre_hoja in libro_original.sheetnames:
            hoja_original = libro_original[nombre_hoja]
            hoja_descargada = libro_descargado[nombre_hoja]

            celdas_originales = _celdas_con_contenido(
                hoja_original
            )

            celdas_descargadas = _celdas_con_contenido(
                hoja_descargada
            )

            for celda_original, celda_descargada in zip_longest(
                celdas_originales,
                celdas_descargadas,
            ):
                if celda_original != celda_descargada:
                    raise RuntimeError(
                        f"{nombre_archivo}: diferencia en "
                        f"la hoja '{nombre_hoja}'. "
                        f"Original: {celda_original!r}. "
                        f"Publicado: {celda_descargada!r}."
                    )

        print(
            f"   ✅ Contenido verificado: {nombre_archivo}"
        )

    finally:
        libro_original.close()

        if libro_descargado is not None:
            libro_descargado.close()


def publicar(sp, drive_id, preparados, originales):
    """
    Publica los Excel preparados y verifica su contenido.

    Si ocurre un fallo, restaura únicamente los archivos
    cuya publicación fue iniciada.

    La restauración se verifica byte por byte contra
    el archivo original.
    """
    archivos_iniciados = []
    archivo_en_proceso = None

    try:
        for nombre, contenido in preparados.items():
            archivo_en_proceso = nombre

            ruta = f"{INPUT_FOLDER}/{nombre}"

            print(f"📤 Publicando {nombre}...")

            # Se registra antes de subir porque una operación
            # fallida podría haber modificado el destino.
            archivos_iniciados.append(nombre)

            sp.upload_file_by_path(
                drive_id=drive_id,
                file_path=ruta,
                file_content=BytesIO(contenido),
            )

            descargado = sp.download_file_by_path(
                drive_id=drive_id,
                file_path=ruta,
            ).getvalue()

            verificar_publicacion_excel(
                original=contenido,
                descargado=descargado,
                nombre_archivo=nombre,
            )

    except Exception as error:
        print("\n⚠️ Error durante la publicación.")
        print(f"Archivo en proceso: {archivo_en_proceso}")
        print(f"Detalle: {error}")
        print("🔄 Restaurando archivos afectados...")

        errores_restauracion = []

        for nombre in reversed(archivos_iniciados):
            ruta = f"{INPUT_FOLDER}/{nombre}"

            try:
                sp.upload_file_by_path(
                    drive_id=drive_id,
                    file_path=ruta,
                    file_content=BytesIO(originales[nombre]),
                )

                restaurado = sp.download_file_by_path(
                    drive_id=drive_id,
                    file_path=ruta,
                ).getvalue()

                if restaurado != originales[nombre]:
                    raise RuntimeError(
                        "La restauración no coincide "
                        "con el archivo original."
                    )

                print(f"   ✅ Restaurado: {nombre}")

            except Exception as exc:
                errores_restauracion.append(
                    f"{nombre}: {exc}"
                )

                print(
                    f"   ❌ Restauración no confirmada: {nombre}"
                )

        if errores_restauracion:
            raise RuntimeError(
                "PUBLICACIÓN FALLIDA. "
                "La restauración quedó incompleta: "
                + "; ".join(errores_restauracion)
            ) from error

        raise RuntimeError(
            "Publicación fallida. Los archivos afectados "
            "fueron restaurados y verificados."
        ) from error

    print("\n✅ Todos los archivos publicados y verificados.")


def ejecutar_reportes() -> None:
    """
    Ejecuta demanda proyectada y auditoría vehicular.

    Cada proceso es independiente: un error de auditoría
    no invalida una actualización de fuentes ya publicada.
    """
    print("\n" + "=" * 65)
    print("🐔 CGO — GENERACIÓN DE REPORTES")
    print("=" * 65)

    errores = []

    # 1. Demanda proyectada.
    print("\n🚀 Ejecutando motor de demanda...")

    try:
        from main import ejecutar_motor

        resultado = ejecutar_motor(
            log_callback=print,
            generar_reporte=True,
        )

        print(
            "✅ Demanda publicada: "
            f"{resultado['nombre_reporte']}"
        )

    except Exception as error:
        errores.append(f"Demanda: {error}")
        print(f"❌ Error en demanda: {error}")

    # 2. Auditoría vehicular.
    print("\n🔍 Ejecutando auditoría de vehículos...")

    ruta_auditoria = (
        Path(__file__).resolve().parents[1]
        / "tests"
        / "auditoria_vehiculos.py"
    )

    try:
        subprocess.run(
            [sys.executable, str(ruta_auditoria)],
            check=True,
        )

        print("✅ Auditoría vehicular completada.")

    except (subprocess.CalledProcessError, OSError) as error:
        errores.append(f"Auditoría vehicular: {error}")
        print(f"❌ Error en auditoría: {error}")

    if errores:
        raise RuntimeError(
            "Las fuentes publicadas no se han revertido. "
            "Hubo errores al generar los reportes:\n"
            + "\n".join(errores)
        )

    print("\n🐔 Proceso integral completado correctamente.")


def ejecutar(publicacion: bool = False):
    inicio = datetime.now(ZoneInfo("America/Lima"))

    print("\n" + "=" * 65)
    print("🐔 CGO — ACTUALIZADOR DE FUENTES")
    print("MODO:", "PUBLICACIÓN" if publicacion else "SIMULACIÓN")
    print("=" * 65)

    sp, drive_id = conectar_sharepoint()

    nombres = [
        TALLER_FILE,
        CLIENTES_FILE,
        SERVICIOS_FILE,
        MODELOS_FILE,
    ]

    originales = {}
    fuentes = {}

    for nombre in nombres:
        print(f"📥 Leyendo {nombre}...")
        originales[nombre], fuentes[nombre] = descargar(
            sp, drive_id, nombre
        )
        print(f"   ✅ {len(fuentes[nombre]):,} filas")

    ordenes = obtener_ordenes(sector="13")

    ids_existentes = set(
        fuentes[TALLER_FILE]["Id doc"].map(texto)
    )

    ids_nuevos = sorted(
        set(ordenes["ID_OPERACION"].map(texto))
        - ids_existentes - {""}
    )

    print(f"\n🐔 Órdenes nuevas: {len(ids_nuevos):,}")

    facturacion = obtener_facturacion(
        ids_ordenes=ids_nuevos,
        sector="13",
    )

    resultado = preparar_fuentes(
        ordenes=ordenes,
        facturacion=facturacion,
        taller_actual=fuentes[TALLER_FILE],
        clientes_actual=fuentes[CLIENTES_FILE],
        servicios_actual=fuentes[SERVICIOS_FILE],
        modelos_actual=fuentes[MODELOS_FILE],
    )

    # El historial de clientes permanece intacto.
    nuevos_clientes = resultado.clientes.iloc[
        len(fuentes[CLIENTES_FILE]):
    ].copy()

    clientes_filtrados = filtrar_clientes_nuevos(
        fuentes[CLIENTES_FILE],
        nuevos_clientes,
    )

    taller_final = resultado.taller

    clientes_final = pd.concat(
        [fuentes[CLIENTES_FILE], clientes_filtrados],
        ignore_index=True,
    )

    validar_incremento(
        fuentes[TALLER_FILE],
        taller_final,
        resultado.nuevas_ordenes,
        TALLER_FILE,
    )

    validar_incremento(
        fuentes[CLIENTES_FILE],
        clientes_final,
        clientes_filtrados,
        CLIENTES_FILE,
    )

    # Los diccionarios con nuevas equivalencias se preparan
    # sin modificar sus filas existentes.
    nuevos_modelos = resultado.diccionario_modelos.iloc[
        len(fuentes[MODELOS_FILE]):
    ].copy()

    nuevas_categorias = resultado.diccionario_servicios.iloc[
        len(fuentes[SERVICIOS_FILE]):
    ].copy()

    if not nuevas_categorias.empty:
        raise RuntimeError(
            "Existen categorías nuevas con jerarquía pendiente "
            "de validar. No se publicarán automáticamente."
        )

    incrementos = {
        TALLER_FILE: resultado.nuevas_ordenes,
        CLIENTES_FILE: clientes_filtrados,
        MODELOS_FILE: nuevos_modelos,
    }

    preparados = {}

    for nombre, nuevas in incrementos.items():
        if nuevas.empty:
            continue

        salida = agregar_filas_excel(
            originales[nombre],
            nuevas,
            nombre,
        )

        contenido = salida.getvalue()

        verificar_archivo(
            contenido,
            len(fuentes[nombre]) + len(nuevas),
            nombre,
        )

        preparados[nombre] = contenido

    print("\n" + "=" * 65)
    print("🐔 RESULTADO DE PREPARACIÓN")
    print("=" * 65)

    print(f"Órdenes incorporables : {len(resultado.nuevas_ordenes):,}")
    print(f"Órdenes pendientes    : {len(resultado.pendientes):,}")
    print(f"Clientes propuestos   : {len(nuevos_clientes):,}")
    print(f"Clientes a incorporar : {len(clientes_filtrados):,}")
    print(f"Modelos nuevos        : {len(nuevos_modelos):,}")

    print("\nArchivos que cambiarían:")
    for nombre in preparados:
        print(f"  - {nombre}")

    if not publicacion:
        print("\n✅ Simulación finalizada. Sin escrituras.")
        return

    if not preparados:
        print("\n✅ No hay archivos que actualizar.")
        print("🐔 Se generarán los reportes con las fuentes actuales.")
        ejecutar_reportes()
        return

    respuesta = input(
        "\n¿Confirmas la publicación en SharePoint? "
        "Escribe PUBLICAR: "
    ).strip()

    if respuesta != "PUBLICAR":
        print("Publicación cancelada.")
        return

    # Verificamos que las fuentes de producción no hayan
    # cambiado desde que fueron descargadas.
    for nombre in preparados:
        actual = sp.download_file_by_path(
            drive_id=drive_id,
            file_path=f"{INPUT_FOLDER}/{nombre}",
        ).getvalue()

        if actual != originales[nombre]:
            raise RuntimeError(
                f"{nombre} cambió durante la ejecución. "
                "Se cancela para evitar sobrescribir cambios."
            )

    # Backups obligatorios de los archivos raíz.
    respaldar(
        sp,
        drive_id,
        originales,
        archivos=list(preparados),
    )

    publicar(
        sp,
        drive_id,
        preparados,
        originales,
    )

    print("\n✅ FUENTES PUBLICADAS CORRECTAMENTE")

    ejecutar_reportes()

    fin = datetime.now(ZoneInfo("America/Lima"))

    print("\n" + "=" * 65)
    print("🐔 ACTUALIZACIÓN INTEGRAL FINALIZADA")
    print("=" * 65)

    print(
        f"Duración total: "
        f"{(fin - inicio).total_seconds():.1f} segundos"
    )


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--publicar",
        action="store_true",
        help="Autoriza solicitar publicación en SharePoint.",
    )

    args = parser.parse_args()

    ejecutar(publicacion=args.publicar)


if __name__ == "__main__":
    main()


from __future__ import annotations

import os
import sys
from collections import defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from databricks import sql
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from m365.config import (
    SITE_NAME,
    DRIVE_NAME,
    INPUT_FOLDER,
    OUTPUT_FOLDER,
    TALLER_FILE,
)
from m365.sharepoint import SharePoint
from m365.excel import read_excel


TABLA_SAP = "catalog_lhcl_prod_gp.gp_adv_ai.gema_maestro_vehiculos"
ARCHIVO_MODELOS = "modelos vehiculos.xlsx"
ARCHIVO_REPORTE = "auditoria_calidad_vehiculos.xlsx"

COLUMNAS_TALLER = ["Id doc", "vin", "placa", "modelo", "año"]
COLUMNAS_MODELOS = ["Material", "Modelo", "Año"]


def limpiar(valor, quitar_espacios=False) -> str:
    if pd.isna(valor):
        return ""

    texto = str(valor).strip().upper()

    if texto.endswith(".0") and texto[:-2].isdigit():
        texto = texto[:-2]

    if quitar_espacios:
        texto = "".join(texto.split())

    return texto


def valores_unicos(serie: pd.Series) -> list[str]:
    return sorted({str(v) for v in serie if str(v)})


def unir_valores(serie: pd.Series) -> str:
    return " | ".join(valores_unicos(serie))


def leer_sharepoint() -> tuple[
    SharePoint, str, pd.DataFrame, pd.DataFrame
]:
    cliente = SharePoint()
    site = cliente.get_site(SITE_NAME)

    if site is None:
        raise RuntimeError("No se encontró el sitio de SharePoint.")

    drive = cliente.get_drive(site["id"], DRIVE_NAME)

    if drive is None:
        raise RuntimeError("No se encontró la biblioteca.")

    drive_id = drive["id"]

    def descargar(nombre: str) -> pd.DataFrame:
        archivo = cliente.download_file_by_path(
            drive_id=drive_id,
            file_path=f"{INPUT_FOLDER}/{nombre}",
        )
        return read_excel(archivo, dtype="string")

    print("📥 Leyendo taller.xlsx...")
    taller = descargar(TALLER_FILE)

    print("📥 Leyendo diccionario de modelos...")
    modelos = descargar(ARCHIVO_MODELOS)

    for nombre, df, columnas in (
        ("taller.xlsx", taller, COLUMNAS_TALLER),
        (ARCHIVO_MODELOS, modelos, COLUMNAS_MODELOS),
    ):
        faltantes = set(columnas) - set(df.columns)
        if faltantes:
            raise ValueError(
                f"{nombre}: faltan columnas {sorted(faltantes)}"
            )

    return cliente, drive_id, taller, modelos


def leer_maestro_sap() -> pd.DataFrame:
    load_dotenv(ROOT / ".env")

    hostname = os.getenv("DATABRICKS_SERVER_HOSTNAME")
    http_path = os.getenv("DATABRICKS_HTTP_PATH")

    if not hostname or not http_path:
        raise ValueError("Faltan variables de Databricks.")

    consulta = f"""
        SELECT
            IDEN_VEHICULO,
            PLACA,
            NUMERO_MATERIAL,
            MODELO
        FROM {TABLA_SAP}
    """

    print("📥 Consultando maestro de vehículos SAP...")

    filas = []
    columnas = []

    with sql.connect(
        server_hostname=hostname,
        http_path=http_path,
        auth_type="databricks-oauth",
    ) as conexion:
        with conexion.cursor() as cursor:
            cursor.execute(consulta)
            columnas = [c[0] for c in cursor.description]

            while True:
                lote = cursor.fetchmany(10000)
                if not lote:
                    break
                filas.extend(tuple(fila) for fila in lote)

    return pd.DataFrame.from_records(
        filas,
        columns=columnas,
    )


def preparar_diccionario(df: pd.DataFrame) -> dict:
    trabajo = df.copy()

    trabajo["clave"] = trabajo["Material"].map(limpiar)
    trabajo["modelo"] = trabajo["Modelo"].map(limpiar)
    trabajo["año"] = trabajo["Año"].map(limpiar)

    diccionario = defaultdict(set)

    for fila in trabajo.itertuples(index=False):
        if fila.clave and fila.modelo and fila.año:
            diccionario[fila.clave].add(
                (fila.modelo, fila.año)
            )

    return dict(diccionario)


def traducir_modelo(
    material: str,
    modelo_sap: str,
    diccionario: dict,
) -> tuple[str, str, str]:
    for codigo in (material, modelo_sap):
        opciones = diccionario.get(codigo, set())

        if len(opciones) == 1:
            modelo, año = next(iter(opciones))
            return modelo, año, "TRADUCIDO"

        if len(opciones) > 1:
            return "", "", "DICCIONARIO_AMBIGUO"

    return "", "", "SIN_TRADUCCION"


def preparar_taller(df: pd.DataFrame) -> pd.DataFrame:
    trabajo = df[COLUMNAS_TALLER].copy()

    for columna in COLUMNAS_TALLER:
        trabajo[columna] = trabajo[columna].map(limpiar)

    trabajo["vin_original"] = trabajo["vin"]
    trabajo["placa_original"] = trabajo["placa"]

    trabajo["vin"] = trabajo["vin"].map(
        lambda x: limpiar(x, quitar_espacios=True)
    )
    trabajo["placa"] = trabajo["placa"].map(
        lambda x: limpiar(x, quitar_espacios=True)
    )

    return trabajo


def preparar_sap(
    df: pd.DataFrame,
    diccionario: dict,
) -> pd.DataFrame:
    trabajo = df.copy()

    for columna in (
        "IDEN_VEHICULO",
        "PLACA",
        "NUMERO_MATERIAL",
        "MODELO",
    ):
        trabajo[columna] = trabajo[columna].map(limpiar)

    trabajo["vin_original"] = trabajo["IDEN_VEHICULO"]
    trabajo["placa_original"] = trabajo["PLACA"]

    trabajo["vin"] = trabajo["IDEN_VEHICULO"].map(
        lambda x: limpiar(x, quitar_espacios=True)
    )
    trabajo["placa"] = trabajo["PLACA"].map(
        lambda x: limpiar(x, quitar_espacios=True)
    )

    traducciones = [
        traducir_modelo(
            fila.NUMERO_MATERIAL,
            fila.MODELO,
            diccionario,
        )
        for fila in trabajo.itertuples(index=False)
    ]

    trabajo[
        ["modelo_traducido", "año_traducido", "estado_traduccion"]
    ] = pd.DataFrame(
        traducciones,
        index=trabajo.index,
    )

    return trabajo


def construir_auditoria(
    taller: pd.DataFrame,
    sap: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    observaciones = []
    vin_observados = set()

    taller_con_vin = taller[taller["vin"] != ""]
    sap_con_vin = sap[sap["vin"] != ""]

    grupos_taller = {
        vin: grupo
        for vin, grupo in taller_con_vin.groupby("vin")
    }

    grupos_sap = {
        vin: grupo
        for vin, grupo in sap_con_vin.groupby("vin")
        if vin in grupos_taller
    }

    # Una placa asociada a distintos VIN es una observación.
    placas_vin = (
        taller_con_vin[taller_con_vin["placa"] != ""]
        .groupby("placa")["vin"]
        .nunique()
    )
    placas_conflictivas = set(
        placas_vin[placas_vin > 1].index
    )

    sap_placas_vin = (
        sap_con_vin[sap_con_vin["placa"] != ""]
        .groupby("placa")["vin"]
        .nunique()
    )
    placas_conflictivas_sap = set(
        sap_placas_vin[sap_placas_vin > 1].index
    )

    for vin, grupo_taller in grupos_taller.items():
        grupo_sap = grupos_sap.get(vin)
        motivos = []

        placas_taller = set(
            valores_unicos(grupo_taller["placa"])
        )
        modelos_taller = set(
            valores_unicos(grupo_taller["modelo"])
        )
        años_taller = set(
            valores_unicos(grupo_taller["año"])
        )

        if len(vin) != 17:
            motivos.append("LONGITUD_VIN")

        if len(placas_taller) > 1:
            motivos.append("VIN_VARIAS_PLACAS_TALLER")

        if len(modelos_taller) > 1:
            motivos.append("VIN_VARIOS_MODELOS_TALLER")

        if len(años_taller) > 1:
            motivos.append("VIN_VARIOS_AÑOS_TALLER")

        if placas_taller & placas_conflictivas:
            motivos.append("PLACA_VARIOS_VIN_TALLER")

        if not placas_taller:
            motivos.append("SIN_PLACA_TALLER")

        if not modelos_taller:
            motivos.append("SIN_MODELO_TALLER")

        if not años_taller:
            motivos.append("SIN_AÑO_TALLER")

        if (
            grupo_taller["vin_original"]
            != grupo_taller["vin"]
        ).any():
            motivos.append("ESPACIOS_EN_VIN_TALLER")

        if (
            grupo_taller["placa_original"]
            != grupo_taller["placa"]
        ).any():
            motivos.append("ESPACIOS_EN_PLACA_TALLER")

        placas_sap = set()
        modelos_sap = set()
        años_sap = set()
        materiales_sap = set()
        codigos_sap = set()
        estado_traduccion = ""

        if grupo_sap is None:
            motivos.append("VIN_NO_ENCONTRADO_SAP")
        else:
            placas_sap = set(
                valores_unicos(grupo_sap["placa"])
            )
            modelos_sap = set(
                valores_unicos(grupo_sap["modelo_traducido"])
            )
            años_sap = set(
                valores_unicos(grupo_sap["año_traducido"])
            )
            materiales_sap = set(
                valores_unicos(grupo_sap["NUMERO_MATERIAL"])
            )
            codigos_sap = set(
                valores_unicos(grupo_sap["MODELO"])
            )

            if len(placas_sap) > 1:
                motivos.append("VIN_VARIAS_PLACAS_SAP")

            if len(modelos_sap) > 1:
                motivos.append("VIN_VARIOS_MODELOS_SAP")

            if len(años_sap) > 1:
                motivos.append("VIN_VARIOS_AÑOS_SAP")

            if placas_sap & placas_conflictivas_sap:
                motivos.append("PLACA_VARIOS_VIN_SAP")

            if placas_taller and placas_sap:
                if placas_taller != placas_sap:
                    motivos.append("DIFERENCIA_PLACA_TALLER_SAP")

            if modelos_taller and modelos_sap:
                if modelos_taller != modelos_sap:
                    motivos.append("DIFERENCIA_MODELO_TALLER_SAP")

            if años_taller and años_sap:
                if años_taller != años_sap:
                    motivos.append("DIFERENCIA_AÑO_TALLER_SAP")

            estados = set(grupo_sap["estado_traduccion"])
            estados.discard("TRADUCIDO")

            if estados:
                motivos.append("TRADUCCION_SAP_PENDIENTE")

            estado_traduccion = " | ".join(
                sorted(set(grupo_sap["estado_traduccion"]))
            )

            if (
                grupo_sap["vin_original"]
                != grupo_sap["vin"]
            ).any():
                motivos.append("ESPACIOS_EN_VIN_SAP")

            if (
                grupo_sap["placa_original"]
                != grupo_sap["placa"]
            ).any():
                motivos.append("ESPACIOS_EN_PLACA_SAP")

        if not motivos:
            continue

        vin_observados.add(vin)

        observaciones.append({
            "vin": vin,
            "motivos": " | ".join(sorted(set(motivos))),
            "cantidad_motivos": len(set(motivos)),
            "ordenes_taller": grupo_taller["Id doc"].nunique(),
            "placas_taller": " | ".join(sorted(placas_taller)),
            "placas_sap": " | ".join(sorted(placas_sap)),
            "modelo_taller": " | ".join(sorted(modelos_taller)),
            "modelo_sap_traducido": " | ".join(
                sorted(modelos_sap)
            ),
            "año_taller": " | ".join(sorted(años_taller)),
            "año_sap_traducido": " | ".join(sorted(años_sap)),
            "materiales_sap": " | ".join(sorted(materiales_sap)),
            "codigos_modelo_sap": " | ".join(sorted(codigos_sap)),
            "estado_traduccion": estado_traduccion,
        })

    # Registros sin VIN: requieren identificación por orden.
    taller_sin_vin = taller[taller["vin"] == ""]

    for fila in taller_sin_vin.itertuples(index=False):
        observaciones.append({
            "vin": "",
            "motivos": "SIN_VIN_TALLER",
            "cantidad_motivos": 1,
            "ordenes_taller": 1,
            "placas_taller": fila.placa,
            "placas_sap": "",
            "modelo_taller": fila.modelo,
            "modelo_sap_traducido": "",
            "año_taller": fila.año,
            "año_sap_traducido": "",
            "materiales_sap": "",
            "codigos_modelo_sap": "",
            "estado_traduccion": "",
        })

    auditoria = pd.DataFrame(observaciones)

    detalle_taller = taller[
        taller["vin"].isin(vin_observados)
        | taller["vin"].eq("")
    ].copy()

    detalle_sap = sap[
        sap["vin"].isin(vin_observados)
    ].copy()

    return auditoria, detalle_taller, detalle_sap


def construir_resumen(
    auditoria: pd.DataFrame,
    total_taller: int,
    total_sap: int,
) -> pd.DataFrame:
    indicadores = [
        ("Registros de taller analizados", total_taller),
        ("Registros del maestro SAP consultados", total_sap),
        ("Filas de auditoría", len(auditoria)),
    ]

    if not auditoria.empty:
        indicadores.append((
            "VIN observados distintos",
            auditoria.loc[
                auditoria["vin"] != "", "vin"
            ].nunique(),
        ))

        motivos = (
            auditoria["motivos"]
            .str.split(" | ", regex=False)
            .explode()
            .value_counts()
        )

        indicadores.extend(
            (f"Motivo: {motivo}", int(cantidad))
            for motivo, cantidad in motivos.items()
        )

    return pd.DataFrame(
        indicadores,
        columns=["indicador", "cantidad"],
    )


def generar_excel(
    resumen: pd.DataFrame,
    auditoria: pd.DataFrame,
    detalle_taller: pd.DataFrame,
    detalle_sap: pd.DataFrame,
) -> BytesIO:
    archivo = BytesIO()

    with pd.ExcelWriter(
        archivo,
        engine="openpyxl",
    ) as writer:
        for nombre, df in (
            ("RESUMEN", resumen),
            ("INCONSISTENCIAS", auditoria),
            ("DETALLE_TALLER", detalle_taller),
            ("DETALLE_SAP", detalle_sap),
        ):
            df.to_excel(
                writer,
                sheet_name=nombre,
                index=False,
            )

            hoja = writer.sheets[nombre]
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions

            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter

            for celda in hoja[1]:
                celda.font = Font(
                    bold=True,
                    color="FFFFFF",
                )
                celda.fill = PatternFill(
                    fill_type="solid",
                    fgColor="203864",
                )

            for indice, columna in enumerate(
                df.columns, start=1
            ):
                ancho = min(
                    max(len(str(columna)) + 3, 17),
                    48,
                )
                hoja.column_dimensions[
                    get_column_letter(indice)
                ].width = ancho

            # Evita notación científica y pérdida de ceros
            # en identificadores almacenados como texto.
            for fila in hoja.iter_rows(
                min_row=2
            ):
                for celda in fila:
                    if isinstance(celda.value, str):
                        celda.number_format = "@"

    archivo.seek(0)
    return archivo


def main() -> None:
    print("\n🐔 AUDITORÍA VEHICULAR — CGO")
    print("=" * 55)

    cliente, drive_id, taller_raw, modelos_raw = (
        leer_sharepoint()
    )

    diccionario = preparar_diccionario(modelos_raw)
    taller = preparar_taller(taller_raw)

    sap_raw = leer_maestro_sap()
    sap = preparar_sap(sap_raw, diccionario)

    print("🔍 Identificando inconsistencias...")

    auditoria, detalle_taller, detalle_sap = (
        construir_auditoria(taller, sap)
    )

    resumen = construir_resumen(
        auditoria,
        len(taller),
        len(sap),
    )

    archivo = generar_excel(
        resumen,
        auditoria,
        detalle_taller,
        detalle_sap,
    )

    ruta = f"{OUTPUT_FOLDER}/{ARCHIVO_REPORTE}"

    print("📤 Publicando auditoría en SharePoint...")

    cliente.upload_file_by_path(
        drive_id=drive_id,
        file_path=ruta,
        file_content=archivo,
    )

    print("\n✅ Auditoría publicada correctamente.")
    print(f"📄 Archivo: {ARCHIVO_REPORTE}")
    print(f"📁 Carpeta: {OUTPUT_FOLDER}")
    print(f"🔎 Filas de auditoría: {len(auditoria):,}")
    print(f"📅 Ejecución: {datetime.now(ZoneInfo('America/Lima')):%Y-%m-%d %H:%M}")


if __name__ == "__main__":
    main()
